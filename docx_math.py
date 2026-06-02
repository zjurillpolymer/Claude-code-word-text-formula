"""LaTeX → Word native OMML formulas via pandoc backend.

Usage:
    from docx_math import add_formula
    from docx import Document

    doc = Document()
    add_formula(doc.add_paragraph(), r'\Delta G = \Delta H - T \Delta S')
    doc.save('output.docx')
"""

import io, os, re, subprocess, tempfile, zipfile
from copy import deepcopy
from lxml import etree

NSW = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NSM = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
W = '{' + NSW + '}'
M = '{' + NSM + '}'
NS = {'w': NSW, 'm': NSM}

_PANDOC_CACHE = {}
_HAS_PANDOC = None


def _check_pandoc():
    global _HAS_PANDOC
    if _HAS_PANDOC is None:
        try:
            subprocess.run(['pandoc', '--version'], capture_output=True)
            _HAS_PANDOC = True
        except FileNotFoundError:
            _HAS_PANDOC = False
    return _HAS_PANDOC


def _convert_via_pandoc(formulas):
    """Convert a list of (latex, display_bool) to OMML elements using pandoc.

    Returns a list of (list_of_omml_elements) matching the input order.
    Each element is a *single* lxml etree Element (m:oMath or m:oMathPara).
    """
    md_lines = []
    for i, (latex, display) in enumerate(formulas):
        if display:
            md_lines.append(f'MARKER{i}\n\n$${latex}$$\n')
        else:
            md_lines.append(f'MARKER{i} ${latex}$ END{i}\n')

    md_text = '\n'.join(md_lines)

    # Run pandoc
    with tempfile.TemporaryDirectory() as tmpdir:
        md_path = os.path.join(tmpdir, 'input.md')
        docx_path = os.path.join(tmpdir, 'output.docx')
        with open(md_path, 'w', encoding='utf-8') as f:
            f.write(md_text)

        r = subprocess.run(
            ['pandoc', md_path, '-f', 'markdown', '-t', 'docx', '-o', docx_path],
            capture_output=True, text=True)
        if r.returncode != 0:
            raise RuntimeError(f'pandoc failed: {r.stderr}')

        # Extract OMML from the generated docx
        with zipfile.ZipFile(docx_path) as z:
            doc_xml = z.read('word/document.xml')

    root = etree.fromstring(doc_xml)
    body = root.find(W + 'body')
    paragraphs = list(body.findall(W + 'p'))

    def para_text(p):
        return ''.join(t.text or '' for t in p.iter(W + 't'))

    results = [None] * len(formulas)

    for i in range(len(paragraphs)):
        p = paragraphs[i]
        text = para_text(p)
        m = re.match(f'MARKER(\\d+)', text)
        if not m:
            continue
        idx = int(m.group(1))

        # Look at current paragraph and next paragraph for OMML
        for cp in (p, paragraphs[i + 1] if i + 1 < len(paragraphs) else None):
            if cp is None:
                continue
            # Check for display math (m:oMathPara)
            paras = cp.findall(M + 'oMathPara')
            if paras:
                results[idx] = [deepcopy(e) for e in paras]
                break
            # Check for inline math (m:oMath)
            maths = cp.findall(M + 'oMath')
            if maths:
                results[idx] = [deepcopy(e) for e in maths]
                break

        if results[idx] is None:
            raise RuntimeError(f'pandoc did not produce OMML for formula {idx}: {formulas[idx][0][:60]}')

    return results


def add_formula(paragraph, latex, display=False):
    """Render a LaTeX formula as a native Word OMML object in *paragraph*.

    The OMML element is placed as a direct child of ``<w:p>`` (matching
    pandoc's output), which is the standard OOXML placement.

    Parameters
    ----------
    paragraph : docx.text.paragraph.Paragraph
        The paragraph to append the formula into.
    latex : str
        LaTeX math expression (without $$ or $ delimiters).
    display : bool, default=False
        If True, uses m:oMathPara (display-style, often centered).
    """
    if not _check_pandoc():
        raise RuntimeError('pandoc is required but not found. Install it: brew install pandoc')

    key = (latex, display)
    if key not in _PANDOC_CACHE:
        _PANDOC_CACHE[key] = _convert_via_pandoc([key])[0]

    omml_elems = _PANDOC_CACHE[key]
    p_elem = paragraph._element
    for elem in omml_elems:
        p_elem.append(deepcopy(elem))


def add_formula_to_document(doc, latex, display=True):
    """Convenience: add a new paragraph containing a formula."""
    p = doc.add_paragraph()
    add_formula(p, latex, display=display)
    return p


# ── test / demo ─────────────────────────────────────────────────────────

if __name__ == '__main__':
    from docx import Document

    doc = Document()
    doc.add_heading('Polymer Physics Formulas', level=1)

    formulas = [
        (r'\Delta G = \Delta H - T \Delta S', False),
        (r'\frac{d[M]}{dt} = k[M]', True),
        (r'R_g^2 = \frac{1}{N} \sum_{i=1}^{N} (\mathbf{r}_i - \mathbf{r}_{CM})^2', True),
        (r'\sigma = E \varepsilon', False),
        (r'G(t) = \frac{k_B T}{V} \sum_{i} \exp(-t / \tau_i)', True),
        (r'\langle R^2 \rangle = N b^2', False),
        (r'[A] = [A]_0 e^{-k t}', False),
        (r'\int_{-\infty}^{\infty} e^{-x^2} dx = \sqrt{\pi}', True),
    ]

    for latex, display in formulas:
        p = doc.add_paragraph()
        add_formula(p, latex, display)

    path = '/Users/arcadio/Documents/formula_test_fixed.docx'
    doc.save(path)
    print(f'Saved to {path}')
